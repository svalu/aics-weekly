# -*- coding: utf-8 -*-
"""
doc/AI Communication Service 업무진행(Weekly).xlsx 에서
9월 현행분을 뽑아 supabase/seed.sql 을 생성한다.

  python scripts/extract.py
"""
import io, json, os, re, unicodedata
from datetime import date, datetime
import warnings
warnings.filterwarnings("ignore")
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "doc", "AI Communication Service 업무진행(Weekly).xlsx")
OUT = os.path.join(ROOT, "supabase", "seed.sql")
YEAR = 2026

# ── 팀 명단 ─────────────────────────────────────────────────
# 실제 사람 이름과 주소는 저장소에 넣지 않는다.
# scripts/roster.local.json (.gitignore 대상) 에 아래 형태로 두면 읽어간다:
#
#   {
#     "order":  ["홍길동", "김철수"],
#     "org":    { "홍길동": ["Sales", "리더"], "김철수": ["AICC", "PM"] },
#     "admins": ["홍길동"],
#     "roman":  { "홍길동": "hong.gildong" },
#     "emails": { "홍길동": "hong@example.com" }
#   }
#
# 파일이 없으면 엑셀의 Owner 칸에서 이름을 모아 쓰고, 파트 · 역할은 비워 둔다.
_LOCAL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "roster.local.json")
_R = json.load(io.open(_LOCAL, encoding="utf-8")) if os.path.exists(_LOCAL) else {}

EMAIL = _R.get("emails", {})
ROMAN = _R.get("roman", {})
ORG = {k: tuple(v) for k, v in _R.get("org", {}).items()}
ADMINS = set(_R.get("admins", []))
ORDER = list(_R.get("order", []))


# ── 유틸 ────────────────────────────────────────────────────
def q(v):
    """SQL 리터럴."""
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, (date, datetime)):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").strip()
    return "'" + s.replace("'", "''") + "'"


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    s = unicodedata.normalize("NFC", str(v)).replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t]+\n", "\n", s)
    return s.strip()


def d(v):
    """셀 → date"""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})", clean(v))
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def email_for(name):
    name = clean(name)
    if name in EMAIL:
        return EMAIL[name], False
    if name in ROMAN:
        return ROMAN[name] + "@mz.co.kr", True
    slug = re.sub(r"[^a-z0-9]+", "", name.lower()) or "member"
    return slug + "@mz.co.kr", True


BLANK_CELL = re.compile(r"^<금주\s*업무>\s*<차주\s*업무>\s*<이슈\s*사항>\s*$")


def split_report(text):
    """'<금주 업무>…<차주 업무>…<이슈 사항>…' → (this, next, issues)"""
    t = clean(text)
    if not t:
        return "", "", ""
    if BLANK_CELL.match(re.sub(r"\n+", " ", t).strip()):
        return "", "", ""
    parts = re.split(r"<\s*(금주\s*업무|차주\s*업무|이슈\s*사항)\s*>", t)
    if len(parts) == 1:
        return t, "", ""
    out = {"금주": "", "차주": "", "이슈": ""}
    i = 1
    while i < len(parts) - 1:
        key = re.sub(r"\s+", "", parts[i])[:2]
        if key in out:
            out[key] = parts[i + 1].strip()
        i += 2
    return out["금주"], out["차주"], out["이슈"]


def parse_week_header(text):
    """'2026. 9월 3주차(09/14~09/18)' → date(2026, 9, 14)"""
    t = clean(text)
    m = re.search(r"\((\d{1,2})/(\d{1,2})\s*~", t)
    if not m:
        return None
    ym = re.match(r"\s*(\d{4})", t)
    return date(int(ym.group(1)) if ym else YEAR, int(m.group(1)), int(m.group(2)))


def monday(dt):
    return dt.fromordinal(dt.toordinal() - dt.weekday()) if dt else None


def parse_meeting_week(text):
    """'8월 03일주' / '9월07일주' → 그 주 월요일"""
    m = re.match(r"(\d{1,2})월(\d{1,2})일?주", re.sub(r"\s+", "", clean(text)))
    return monday(date(YEAR, int(m.group(1)), int(m.group(2)))) if m else None


KIND_RE = re.compile(r"^\s*\[([^\]]+)\]\s*")
PERIOD_RE = re.compile(
    r"\((\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})\s*~\s*(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})\)")


def parse_project_cell(text):
    """'[프로젝트] AIR Connect 개발 P1-2\n(2026.07.06 ~ 2026.12.31)' 파싱."""
    t = clean(text)
    if not t:
        return None
    kind = "프로젝트"
    m = KIND_RE.match(t)
    if m:
        kind = m.group(1).strip()
        t = KIND_RE.sub("", t)
    sd = ed = None
    pm = PERIOD_RE.search(t)
    if pm:
        g = [int(x) for x in pm.groups()]
        sd, ed = date(g[0], g[1], g[2]), date(g[3], g[4], g[5])
        t = PERIOD_RE.sub("", t)
    t = re.sub(r"\(\s*\?\s*\)", "", t)
    return {"name": re.sub(r"\s*\n\s*", " ", t).strip(" \n-"),
            "kind": kind, "start": sd, "end": ed}


def collect_names(wb):
    """roster.local.json 이 없을 때, 엑셀 Owner 칸에서 이름을 모은다."""
    seen = []
    for sheet, cols in [("Action Tracker", (5,)), ("Sales Activity", (2,))]:
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(values_only=True):
            for c in cols:
                if c >= len(r):
                    continue
                n = clean(r[c])
                if (2 <= len(n) <= 4 and n not in ("Owner", "담당", "담당자")
                        and all("가" <= ch <= "힣" for ch in n)
                        and n not in seen):
                    seen.append(n)
    return seen


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    lines = []
    W = lines.append

    global ORDER
    if not ORDER:
        ORDER = collect_names(wb)
        print("roster.local.json 이 없어 엑셀에서 이름 %d명을 모았습니다." % len(ORDER))

    W("-- ============================================================")
    W("--  AICS Weekly · 시드 데이터 (자동 생성 · scripts/extract.py)")
    W("--  출처: doc/AI Communication Service 업무진행(Weekly).xlsx")
    W("--  생성: " + datetime.now().strftime("%Y-%m-%d %H:%M"))
    W("--  schema.sql 을 먼저 실행한 뒤 이 파일을 실행하세요.")
    W("-- ============================================================")
    W("begin;")
    W("")

    # ── 팀원 ───────────────────────────────────────────────
    todo = []
    W("-- ---------- 팀원 ----------")
    W("-- 주의: TODO 표시된 이메일은 추측값입니다. 실제 주소로 바꾼 뒤 실행하세요.")
    for i, n in enumerate(ORDER):
        mail, guessed = email_for(n)
        part, role = ORG.get(n, (None, None))
        if guessed:
            todo.append((n, mail))
        W("insert into members (email, name, part, role, is_admin, sort_order) values "
          "({}, {}, {}, {}, {}, {}) on conflict (email) do nothing;{}".format(
              q(mail), q(n), q(part), q(role), q(n in ADMINS), (i + 1) * 10,
              "   -- TODO: 실제 이메일" if guessed else ""))
    W("")

    # ── Ground Rule ────────────────────────────────────────
    W("-- ---------- Ground Rule ----------")
    order = 0
    for row in wb["Ground Rule"].iter_rows(values_only=True):
        for cell in row:
            t = clean(cell)
            if t and t != "Ground Rule" and len(t) > 10:
                order += 10
                W("insert into ground_rules (sort_order, body) values ({}, {});".format(order, q(t)))
                break
    W("")

    # ── 프로젝트 (9월 Tech 시트) ───────────────────────────
    W("-- ---------- 프로젝트 ----------")
    rows = list(wb["Weekly 주간회의 9월(Tech)"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if clean(r[0]).startswith("프로젝트(업무)"))
    projects = {}
    order = 0
    for r in rows[hdr + 1:]:
        info = parse_project_cell(r[0])
        if not info or not info["name"] or info["name"].startswith("📋"):
            continue
        if info["name"] in projects:
            continue
        order += 10
        info.update(owner=clean(r[1]) or None,
                    staff=[s.strip() for s in clean(r[2]).split("\n") if s.strip()],
                    order=order)
        projects[info["name"]] = info
        W("insert into projects (name, kind, owner_id, start_date, end_date, sort_order) "
          "select {}, {}, (select id from members where name = {}), {}, {}, {} "
          "where not exists (select 1 from projects where name = {});".format(
              q(info["name"]), q(info["kind"]), q(info["owner"]),
              q(info["start"]), q(info["end"]), order, q(info["name"])))
    W("")

    W("-- ---------- 프로젝트 투입인력 ----------")
    for info in projects.values():
        for person in sorted(set(info["staff"] + ([info["owner"]] if info["owner"] else []))):
            if person not in ORDER:
                continue
            W("insert into project_members (project_id, member_id) "
              "select p.id, m.id from projects p, members m "
              "where p.name = {} and m.name = {} on conflict do nothing;".format(
                  q(info["name"]), q(person)))
    W("")

    # ── 주간보고 (9월 개인_Tech 시트) ──────────────────────
    W("-- ---------- 주간보고 ----------")
    rows = list(wb["Weekly 주간회의 9월(개인_Tech)"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if any("주차" in clean(c) for c in r))
    weeks = [parse_week_header(c) for c in rows[hdr]]
    n_reports = 0
    for r in rows[hdr + 1:]:
        head_lines = [x.strip() for x in clean(r[0]).split("\n") if x.strip()]
        if not head_lines or head_lines[0] not in ORDER:
            continue
        person = head_lines[0]
        pname = None
        if len(head_lines) > 1:
            proj = parse_project_cell("\n".join(head_lines[1:]))
            pname = proj["name"] if proj and proj["name"] else None
            if pname:  # Tech 시트의 정식 프로젝트명으로 맞춘다
                pname = next((k for k in projects if pname in k or k in pname), pname)
        for ci, cell in enumerate(r):
            if ci == 0 or ci >= len(weeks) or weeks[ci] is None:
                continue
            tw, nw, iss = split_report(cell)
            if not (tw or nw or iss):
                continue
            n_reports += 1
            proj_sql = "(select id from projects where name = {})".format(q(pname)) if pname else "null"
            W("insert into weekly_reports (member_id, project_id, week_start, this_week, next_week, issues) "
              "select m.id, {}, {}, {}, {}, {} from members m where m.name = {};".format(
                  proj_sql, q(weeks[ci]), q(tw), q(nw), q(iss), q(person)))
    W("")

    # ── Action Tracker ─────────────────────────────────────
    W("-- ---------- Action Tracker ----------")
    rows = list(wb["Action Tracker"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows)
               if clean(r[1]) == "#" and "Owner" in [clean(c) for c in r])
    cols = {clean(c): i for i, c in enumerate(rows[hdr]) if clean(c)}
    n_actions = 0
    for r in rows[hdr + 1:]:
        title = clean(r[cols["Action Item"]])
        if not title:
            continue
        seq = clean(r[cols["#"]])
        status = clean(r[cols["Status"]]) or "20.진행중"
        if status == "30.완료":
            status = "50.Close"
        n_actions += 1
        W("insert into action_items (seq, category, customer, title, owner_id, status, "
          "progress_update, open_date, target_date, start_date, close_date, remark) values "
          "({}, {}, {}, {}, (select id from members where name = {}), {}, {}, {}, {}, {}, {}, {});".format(
              seq if seq.isdigit() else "null",
              q(clean(r[cols["구분"]]) or "Sales"),
              q(clean(r[cols["고객명"]])),
              q(title),
              q(clean(r[cols["Owner"]])),
              q(status),
              q(clean(r[cols["Progress Update"]])),
              q(d(r[cols["Open Date"]])), q(d(r[cols["Target Date"]])),
              q(d(r[cols["Start Date"]])), q(d(r[cols["Close Date"]])),
              q(clean(r[cols["Remark"]])) if "Remark" in cols else "''"))
    W("")

    # ── Sales 파이프라인 ───────────────────────────────────
    W("-- ---------- Sales 파이프라인 ----------")
    rows = list(wb["Weekly 주간회의(Sales)"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if clean(r[0]) == "No.")
    n_deals = 0
    for r in rows[hdr + 1:]:
        no, name = clean(r[0]), clean(r[3])
        if not no or not name:
            continue
        n_deals += 1
        start, end = clean(r[8]), clean(r[9])
        period = (start + (" ~ " + end if end else "")) if start else ""
        notes = clean(r[12]) if len(r) > 12 else ""
        W("with dl as (insert into deals (seq, category, stage, name, customer, lead, "
          "sales_aws, sales_mzc, period, amount, seats, remark) values "
          "({}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {}) returning id)".format(
              no.replace(".0", "") if re.match(r"^\d+(\.0)?$", no) else "null",
              q(clean(r[1]) or "Sales"), q(clean(r[2]) or "진행중"), q(name),
              q(clean(r[4])), q(clean(r[5]).replace("\n", ", ")),
              q(clean(r[6])), q(clean(r[7]).replace("\n", ", ")),
              q(period), q(clean(r[10])), q(clean(r[11])),
              q(clean(r[13]) if len(r) > 13 else "")))
        blocks = re.split(r"\[(\d{8})\]", notes)
        pairs = []
        if len(blocks) > 1:
            i = 1
            while i < len(blocks) - 1:
                body = blocks[i + 1].strip()
                try:
                    ld = datetime.strptime(blocks[i], "%Y%m%d").date()
                except ValueError:
                    ld = None
                if body:
                    pairs.append((ld, body))
                i += 2
        elif notes:
            pairs.append((None, notes))
        if pairs:
            W("insert into deal_updates (deal_id, log_date, body) values\n  " +
              ",\n  ".join("((select id from dl), {}, {})".format(q(ld), q(b)) for ld, b in pairs) + ";")
        else:
            W("select * from dl;")
    W("")

    # ── 고객 미팅 로그 ─────────────────────────────────────
    W("-- ---------- 고객 미팅 로그 ----------")
    rows = list(wb["Sales Activity"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows)
               if clean(r[1]) == "#" and "Owner" in [clean(c) for c in r])
    n_meet = 0
    for r in rows[hdr + 1:]:
        customer = clean(r[5])
        if not customer:
            continue
        seq = clean(r[1])
        n_meet += 1
        W("insert into meetings (seq, owner_id, week_start, kind, customer, contact, "
          "planned_date, purpose, done, actual_date, companions, notes) values "
          "({}, (select id from members where name = {}), {}, {}, {}, {}, {}, {}, {}, {}, {}, {});".format(
              seq if seq.isdigit() else "null",
              q(clean(r[2])), q(parse_meeting_week(r[3])),
              q(clean(r[4]) or "Net New"), q(customer), q(clean(r[6])),
              q(d(r[7])), q(clean(r[8])), q(clean(r[9])), q(d(r[10])),
              q(clean(r[11]) if len(r) > 11 else ""),
              q(clean(r[12]) if len(r) > 12 else "")))
    W("")
    W("commit;")
    W("")

    io.open(OUT, "w", encoding="utf-8").write("\n".join(lines) + "\n")

    print("생성: " + OUT)
    print("  팀원        %d" % len(ORDER))
    print("  프로젝트    %d" % len(projects))
    print("  주간보고    %d" % n_reports)
    print("  액션아이템  %d" % n_actions)
    print("  딜          %d" % n_deals)
    print("  미팅        %d" % n_meet)
    if todo:
        print("\n  실제 이메일 확인 필요 %d명:" % len(todo))
        for n, m in todo:
            print("     %s -> %s" % (n, m))


if __name__ == "__main__":
    main()

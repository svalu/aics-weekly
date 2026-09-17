# -*- coding: utf-8 -*-
"""
저장소에 올라갈 파일에 실제 회사 데이터가 남아 있는지 검사한다.

원본 엑셀에서 고객사명 · 담당자 실명 · 금액을 뽑아낸 뒤,
git 이 추적할 파일들(= .gitignore 에 걸리지 않는 파일) 안에서 찾는다.

  python scripts/leak_check.py

하나라도 나오면 종료 코드 1.
"""
import io, os, re, subprocess, sys, unicodedata
import warnings
warnings.filterwarnings("ignore")
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "doc", "AI Communication Service 업무진행(Weekly).xlsx")

# 팀원 이름은 공개하기로 한 것이라 검사에서 제외한다
TEAM = {
    "강경표", "민경윤", "이민우", "이진수", "류해원", "이기원", "전병수",
    "이현행", "오대성", "최창현", "김성우", "정택수", "이상묵", "정종호",
    "방희태", "장유석", "김진규", "구창현", "최진용", "이예원",
}
# 일반 명사·머리글자라 고객사 식별에 쓸 수 없는 것들
STOP = {
    "구분", "고객명", "고객사", "고객 담당자", "담당", "담당자", "진행", "완료",
    "예정", "미팅", "내용", "일정", "협의", "확인", "요청", "검토", "보고",
    "작성", "전달", "견적 산출을 위한 정보 수집 중",
    "Sales", "Tech", "ALL", "Owner", "Status", "Internal", "AICC", "IPCC",
    "AWS", "MZC", "AI", "CRM", "SI", "PS", "BPO", "GTM", "POC", "RFP",
    "Net New", "Existing", "N/A", "TBH", "MSP", "SaaS", "SSU", "None",
    "Demo", "https:", "Sales (AWS)", "Sales (MZC)",
}
# 생성물이라 검사 의미가 없는 파일
SKIP_FILES = {"package-lock.json", "next-env.d.ts", "scripts/leak_check.py"}


def clean(v):
    if v is None:
        return ""
    return unicodedata.normalize("NFC", str(v)).strip()


def collect_secrets():
    """엑셀에서 '이게 나오면 유출' 인 문자열을 모은다."""
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    out = set()

    def add(s):
        s = clean(s)
        if not s or s in STOP or s in TEAM:
            return
        if len(s) < 2 or len(s) > 40:
            return
        if re.fullmatch(r"[\d\s.,:/()~-]+", s):
            return
        out.add(s)

    # 고객사 · 담당자
    for sheet, cols in [
        ("Sales Activity", (5, 6)),          # 고객사, 고객 담당자
        ("Action Tracker", (3,)),            # 고객명
        ("Weekly 주간회의(Sales)", (4, 6, 7)),  # 고객명, AWS, MZC
    ]:
        for r in wb[sheet].iter_rows(values_only=True):
            for c in cols:
                if c < len(r):
                    for piece in re.split(r"[,\n/]", clean(r[c])):
                        add(piece)

    # 금액 표기
    for r in wb["Weekly 주간회의(Sales)"].iter_rows(values_only=True):
        if len(r) > 10:
            for piece in re.split(r"\n", clean(r[10])):
                add(piece)

    return {s for s in out if s}


def tracked_files():
    """git 이 추적할 파일 목록 (.gitignore 반영)."""
    try:
        res = subprocess.run(
            ["git", "ls-files", "--cached", "--others", "--exclude-standard"],
            cwd=ROOT, capture_output=True, text=True, encoding="utf-8", check=True)
        files = [f.strip() for f in res.stdout.splitlines() if f.strip()]
        if files:
            return files
    except Exception:
        pass
    # git 저장소가 아직 없을 때의 대략적인 목록
    skip_dirs = {"node_modules", ".next", ".git", ".localdb", "doc", "out"}
    files = []
    for base, dirs, names in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in skip_dirs]
        for n in names:
            p = os.path.relpath(os.path.join(base, n), ROOT).replace("\\", "/")
            if p in ("supabase/seed.sql", "scripts/emails.local.json"):
                continue
            if p.startswith(".env"):
                continue
            files.append(p)
    return files


TEXT_EXT = {".ts", ".tsx", ".js", ".mjs", ".json", ".sql", ".md", ".py",
            ".css", ".html", ".txt", ".yml", ".yaml"}


def main():
    if not os.path.exists(XLSX):
        print("원본 엑셀이 없어 검사할 기준을 만들 수 없습니다:", XLSX)
        return 2

    secrets = collect_secrets()
    print("검사 기준 %d개 (고객사 · 담당자 실명 · 금액)" % len(secrets))

    files = tracked_files()
    print("검사 대상 %d개 파일\n" % len(files))

    hits = []
    for rel in files:
        if rel in SKIP_FILES:
            continue
        ext = os.path.splitext(rel)[1].lower()
        path = os.path.join(ROOT, rel)
        if not os.path.isfile(path):
            continue
        if ext not in TEXT_EXT:
            # 텍스트가 아닌 파일이 딸려 들어오면 그 자체로 알려준다
            if ext in (".xlsx", ".xls", ".csv", ".pdf"):
                hits.append((rel, "(데이터 파일이 그대로 포함됨)"))
            continue
        try:
            body = io.open(path, encoding="utf-8", errors="ignore").read()
        except OSError:
            continue
        for s in secrets:
            if s in body:
                hits.append((rel, s))

    if hits:
        print("유출 발견 %d건:" % len(hits))
        for rel, s in sorted(set(hits))[:60]:
            print("  %-36s  %s" % (rel, s))
        if len(set(hits)) > 60:
            print("  ... 외 %d건" % (len(set(hits)) - 60))
        return 1

    print("깨끗합니다. 커밋될 파일에 실제 고객사 · 담당자 · 금액이 없습니다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
